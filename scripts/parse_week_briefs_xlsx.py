#!/usr/bin/env python3
"""Parse Instructor_Context_Briefs_by_Week.xlsx into week-context-briefs-by-week.json."""

import json
import re
from pathlib import Path

import openpyxl

XLSX = Path(r"c:\Users\LENOVO\Downloads\Instructor_Context_Briefs_by_Week.xlsx")
OUT = Path(__file__).resolve().parent.parent / "src" / "content" / "week-context-briefs-by-week.json"

TRACK_MAP = {"SWE": "swe", "EM": "em", "PM-TPM": "pm"}


def title_from_content(content: str) -> str:
    first = content.strip().split("\n")[0].strip()
    if "—" in first:
        return first.split("—", 1)[-1].strip()
    return first[:120]


def extract_section(content: str, section_num: int) -> str | None:
    pattern = rf"{section_num}\.\s+[^\n]+\n(.*?)(?=\n\d+\.\s+|\Z)"
    m = re.search(pattern, content, re.DOTALL)
    return m.group(1).strip() if m else None


def extract_bullets(text: str | None) -> list[str]:
    if not text:
        return []
    items = []
    for line in text.split("\n"):
        line = line.strip()
        if line.startswith("•") or line.startswith("-"):
            items.append(line.lstrip("•- ").strip())
    return items[:6]


def main() -> None:
    wb = openpyxl.load_workbook(str(XLSX), read_only=True)
    result: dict = {}

    for sheet_name in wb.sheetnames:
        track = TRACK_MAP.get(sheet_name)
        if not track:
            continue
        ws = wb[sheet_name]
        weeks: dict = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            week_label = str(row[0]).strip()
            content = str(row[1] or "").strip()
            if not content:
                continue
            teaching = extract_section(content, 7)
            questions = extract_section(content, 6)
            recap = extract_section(content, 9)
            misconceptions = extract_section(content, 8)
            weeks[week_label] = {
                "weekLabel": week_label,
                "title": title_from_content(content),
                "content": content,
                "oneMinuteRecap": recap or "",
                "teachingEmphasis": extract_bullets(teaching),
                "likelyQuestions": extract_bullets(questions),
                "misconceptions": extract_bullets(misconceptions),
            }
        result[track] = weeks

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {OUT} ({OUT.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
