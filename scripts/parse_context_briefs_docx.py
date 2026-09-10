#!/usr/bin/env python3
"""Parse Instructor Context Briefs docx into instructor-context-briefs.json."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document

DOCX_PATH = Path(r"c:\Users\LENOVO\Downloads\Instructor_Context_Briefs_Agentic_AI_2_0_v2.docx")
OUT_JSON = Path(__file__).resolve().parent.parent / "src" / "content" / "instructor-context-briefs.json"
OUT_INDEX = Path(__file__).resolve().parent.parent / "src" / "content" / "context-brief-module-index.json"

SECTION_ORDER = [
    ("courseJourney", r"^1\.\s*Course Journey\s*(.*)"),
    ("priorKnowledge", r"^2\.\s*What learners already know\s*(.*)"),
    ("learnToday", r"^3\.\s*What learners will learn today\s*(.*)"),
    ("preparesFor", r"^4\.\s*What this prepares learners for\s*(.*)"),
    ("connections", r"^5\.\s*How instructors should connect today's lesson\s*(.*)"),
    ("likelyQuestions", r"^6\.\s*Questions learners are likely to ask\s*(.*)"),
    ("teachingEmphasis", r"^7\.\s*Teaching emphasis\s*(.*)"),
    ("misconceptions", r"^8\.\s*Common misconceptions\s*(.*)"),
    ("oneMinuteRecap", r"^9\.\s*One-minute recap for instructors\s*(.*)"),
]

TRACK_PREFIX = {"SWE": "swe", "EM": "em", "PM": "pm"}

MODULE_HEADER_RE = re.compile(
    r"^(SWE|EM|PM)\s+(?:Module(?:s)?|Capstone)\s+(.+?)\s*[—\-]\s+(.+)$",
    re.IGNORECASE,
)
EM_INTERVIEW_BLOCK_RE = re.compile(r"^EM Interview-Prep Deep Dives", re.I)

WEEK_TOKEN_RE = re.compile(
    r"^(?:W)?(\d+)(?:\(a\)|\(A\)|a|A)?(?:\s*[—\-]\s*)?",
    re.I,
)


def normalize_week_key(raw: str) -> str | None:
    text = raw.strip()
    if not text:
        return None
    lower = text.lower().replace("–", "-").replace("—", "-")
    if "capstone" in lower:
        return "capstone"
    if "orientation" in lower or lower in {"0", "module 0"}:
        return "orientation"
    m = re.search(r"(\d+)\s*\(\s*a\s*\)", lower, re.I)
    if m:
        return f"W{m.group(1)}a"
    m = re.search(r"(\d+)\s*a\b", lower, re.I)
    if m and ("w" in lower or "(" in lower):
        return f"W{m.group(1)}a"
    m = re.search(r"w\s*(\d+)\s*a", lower, re.I)
    if m:
        return f"W{m.group(1)}a"
    m = re.search(r"\(w\s*(\d+)\)", lower, re.I)
    if m:
        return f"W{m.group(1)}"
    m = re.search(r"\bw\s*(\d+)\b", lower, re.I)
    if m:
        return f"W{m.group(1)}"
    m = re.search(r"\b(\d+)\s*\(\s*a\s*\)", lower, re.I)
    if m:
        return f"W{m.group(1)}a"
    m = re.search(r"^(\d+)\s*\(\s*a\s*\)$", lower, re.I)
    if m:
        return f"W{m.group(1)}a"
    m = re.search(r"^(\d+)$", lower)
    if m:
        return f"W{m.group(1)}"
    return None


def infer_week_key(track: str, number_part: str, title: str) -> str:
    combined = f"{number_part} {title}"
    if "capstone" in combined.lower():
        return "capstone"
    if "module 0" in combined.lower() or number_part.strip() == "0":
        return "orientation"
    wk = normalize_week_key(number_part)
    if wk:
        return wk
    m = re.search(r"\(W(\d+(?:\(a\))?)\)", number_part, re.I)
    if m:
        return normalize_week_key(m.group(1)) or normalize_week_key(f"W{m.group(1)}") or "orientation"
    m = re.search(r"\(W(\d+)\s*[-–]\s*(\d+)\)", number_part, re.I)
    if m:
        return "capstone"
    if re.search(r"w\d+\s*[-–]\s*w\d+", lower):
        return "capstone"
    return normalize_week_key(title) or "orientation"


def module_id_for(track: str, week_key: str) -> str:
    prefix = TRACK_PREFIX[track]
    if week_key == "orientation":
        return f"{prefix}-orientation"
    if week_key == "capstone":
        return f"{prefix}-capstone"
    if week_key.startswith("W"):
        suffix = week_key[1:]
        if suffix.endswith("a"):
            core = suffix[:-1].zfill(2)
            return f"{prefix}-w{core}a"
        return f"{prefix}-w{suffix.zfill(2)}"
    return f"{prefix}-{week_key.lower()}"


def parse_misconceptions(lines: list[str]) -> list[dict]:
    items = []
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if '"' in line and "Clarify" in line:
            parts = line.split('" Clarify', 1)
            myth = parts[0].strip().strip('"')
            clarification = "Clarify" + parts[1] if len(parts) > 1 else ""
            items.append({"myth": myth, "clarification": clarification})
        elif "." in line:
            dot = line.find(".")
            items.append({"myth": line[: dot + 1].strip(), "clarification": line[dot + 1 :].strip()})
        else:
            items.append({"myth": line, "clarification": ""})
    return items


def extract_week_refs(text: str) -> list[str]:
    refs: list[str] = []
    for m in re.finditer(r"W(\d+)(?:\(a\)|a)?", text, re.I):
        suffix = "a" if re.search(r"W\d+(?:\(a\)|a)", m.group(0), re.I) and re.search(
            r"(?:\(a\)|a)", m.group(0), re.I
        ) else ""
        key = f"W{m.group(1)}{suffix}"
        if key not in refs:
            refs.append(key)
    return refs


def parse_concept_line(line: str) -> dict | None:
    line = line.strip()
    if not line or line.lower().startswith("concepts and skills"):
        return None
    m = re.match(r"^(W[\d\(\)aA–\-,\s]+)\s*[—\-]\s*(.+)$", line)
    if not m:
        return None
    label_part, text = m.group(1).strip(), m.group(2).strip()
    weeks = extract_week_refs(label_part)
    label = label_part.replace("—", "-").split("-")[0].strip()
    if "–" in label_part or "," in label_part:
        label = re.sub(r"\s+", " ", label_part.replace("–", "–"))
    return {"label": label, "weeks": weeks, "text": text}


def project_kind_from_label(label: str) -> str:
    lower = label.lower()
    if "post-class" in lower:
        return "post-class"
    if "assignment" in lower:
        return "assignment"
    if "optional" in lower:
        return "optional"
    if "interview" in lower:
        return "interview-case"
    if "capstone option" in lower:
        return "capstone-option"
    if "capstone" in lower:
        return "capstone"
    if "rollup" in lower or "earlier builds" in lower:
        return "rollup"
    return "project"


def parse_project_line(line: str) -> dict | None:
    line = line.strip()
    if not line:
        return None
    if line.lower().startswith("projects and assignments"):
        return None
    if line.lower().startswith("prior builds"):
        return {
            "label": "Earlier builds",
            "name": line.split(":", 1)[-1].strip() if ":" in line else line,
            "weeks": extract_week_refs(line),
            "kind": "rollup",
            "text": line,
        }
    if line.lower().startswith("earlier builds"):
        return {
            "label": "Earlier builds",
            "name": line.split(":", 1)[-1].strip() if ":" in line else line,
            "weeks": extract_week_refs(line),
            "kind": "rollup",
            "text": line,
        }
    m = re.match(r"^(W[\d\(\)aA–\-,\s]+)\s+(Project|Assignment|Post-class|Optional)\s*[—\-]\s*(.+)$", line, re.I)
    if m:
        weeks = extract_week_refs(m.group(1))
        kind = project_kind_from_label(m.group(2))
        rest = m.group(3).strip()
        name = rest.split(":", 1)[0].strip() if ":" in rest else rest.split(".")[0].strip()
        return {
            "label": m.group(2),
            "name": name,
            "weeks": weeks,
            "kind": kind,
            "text": rest,
        }
    m = re.match(r"^(W[\d\(\)aA–\-,\s]+)\s+(.+)$", line)
    if m and ("project" in line.lower() or "assignment" in line.lower() or "—" in line):
        weeks = extract_week_refs(m.group(1))
        rest = m.group(2).strip()
        kind = project_kind_from_label(rest)
        name = rest.split(":", 1)[0].strip()
        return {"label": "Build", "name": name, "weeks": weeks, "kind": kind, "text": rest}
    return None


def parse_prior_knowledge(body_lines: list[str]) -> dict:
    concepts: list[dict] = []
    projects: list[dict] = []
    intro: str | None = None
    mode = None
    for line in body_lines:
        stripped = line.strip()
        if not stripped:
            continue
        lower = stripped.lower()
        if lower.startswith("concepts and skills"):
            mode = "concepts"
            continue
        if lower.startswith("projects and assignments"):
            mode = "projects"
            continue
        if lower.startswith("prior builds"):
            mode = "projects"
            proj = parse_project_line(stripped)
            if proj:
                projects.append(proj)
            continue
        if mode == "concepts":
            concept = parse_concept_line(stripped)
            if concept:
                concepts.append(concept)
            elif not concepts and intro is None:
                intro = stripped
            elif concepts:
                concepts[-1]["text"] += " " + stripped
        elif mode == "projects":
            if lower.startswith("earlier builds"):
                proj = parse_project_line(stripped)
                if proj:
                    projects.append(proj)
                continue
            proj = parse_project_line(stripped)
            if proj:
                projects.append(proj)
            elif projects and projects[-1]["kind"] != "rollup":
                projects[-1]["text"] += " " + stripped
        elif intro is None and not lower.startswith("the full track"):
            intro = stripped
        else:
            concept = parse_concept_line(stripped)
            if concept:
                mode = "concepts"
                concepts.append(concept)

    return {
        "intro": intro,
        "conceptsLabel": "Concepts and skills in hand",
        "concepts": concepts,
        "projectsLabel": "Projects and assignments already built",
        "projects": projects,
    }


def parse_brief_sections(lines: list[str]) -> dict:
    sections: dict = {}
    current_key = None
    current_inline = ""
    current_body: list[str] = []
    list_keys = {"connections", "likelyQuestions", "teachingEmphasis", "misconceptions"}

    def flush():
        nonlocal current_key, current_inline, current_body
        if current_key is None:
            return
        if current_key == "priorKnowledge":
            sections[current_key] = parse_prior_knowledge(current_body)
        elif current_key in list_keys:
            if current_key == "misconceptions":
                sections[current_key] = parse_misconceptions(
                    [current_inline] + current_body if current_inline else current_body
                )
            else:
                items = []
                if current_inline:
                    items.append(current_inline)
                for b in current_body:
                    b = b.strip()
                    if b:
                        items.append(b)
                sections[current_key] = items
        else:
            text = current_inline
            if current_body:
                text = (text + " " + " ".join(current_body)).strip() if text else " ".join(current_body).strip()
            sections[current_key] = text
        current_key = None
        current_inline = ""
        current_body = []

    for line in lines:
        matched = False
        for key, pattern in SECTION_ORDER:
            m = re.match(pattern, line.strip(), re.I)
            if m:
                flush()
                current_key = key
                current_inline = m.group(1).strip()
                matched = True
                break
        if not matched and current_key:
            current_body.append(line)

    flush()
    return sections


def parse_orientation_notes(lines: list[str]) -> list[dict]:
    notes = []
    for line in lines:
        t = line.strip()
        if not t or t.lower().startswith("this layer"):
            notes.append({"label": None, "text": t}) if t else None
            continue
        if t.lower().startswith("what it covers"):
            notes.append({"label": "What it covers", "text": t.split(":", 1)[-1].strip()})
        else:
            notes.append({"label": None, "text": t})
    return [n for n in notes if n["text"]]


def find_modules(paragraphs: list[str]) -> list[dict]:
    headers: list[tuple[int, str, str, str, str]] = []
    for i, text in enumerate(paragraphs):
        t = text.strip()
        if not t:
            continue
        if t.startswith("TRACK "):
            continue
        m = MODULE_HEADER_RE.match(t)
        if m:
            track, number_part, title = m.group(1).upper(), m.group(2).strip(), m.group(3).strip()
            if track == "PM" and "TPM" in t:
                track = "PM"
            headers.append((i, track, number_part, title, t))
            continue
        if t.startswith("EM Capstone"):
            headers.append((i, "EM", "capstone", "capstone", t))
        if t.startswith("PM Capstone"):
            headers.append((i, "PM", "capstone", "capstone", t))

    modules = []
    for idx, (start, track, number_part, title, raw) in enumerate(headers):
        end = headers[idx + 1][0] if idx + 1 < len(headers) else len(paragraphs)
        body = [p for p in paragraphs[start + 1 : end] if p.strip()]
        if "capstone" in raw.lower():
            week_key = "capstone"
        else:
            week_key = infer_week_key(track, number_part, title)
        mod_id = module_id_for(track, week_key)
        kind = "orientation" if week_key == "orientation" else "brief"
        if "(Shared framing module" in " ".join(body[:3]):
            kind = "shared"

        module = {
            "id": mod_id,
            "track": TRACK_PREFIX[track],
            "order": len(modules) + 1,
            "code": raw.split("—")[0].strip() if "—" in raw else raw,
            "title": title,
            "subtitle": None,
            "weekKey": week_key,
            "weekLabel": f"Week {week_key[1:]}" if week_key.startswith("W") else week_key.title(),
            "weeks": [week_key] if week_key not in {"orientation", "capstone"} else [],
            "phase": phase_for(week_key, title),
            "kind": kind,
        }

        if kind == "orientation":
            module["notes"] = parse_orientation_notes(body)
        elif kind == "shared":
            module["notes"] = parse_shared_notes(body)
            sw = shared_with_for(mod_id)
            if sw:
                module["sharedWith"] = sw
        else:
            sections = parse_brief_sections(body)
            if sections:
                module["sections"] = {
                    "courseJourney": sections.get("courseJourney", ""),
                    "priorKnowledge": sections.get(
                        "priorKnowledge",
                        {
                            "intro": None,
                            "conceptsLabel": "Concepts and skills in hand",
                            "concepts": [],
                            "projectsLabel": "Projects and assignments already built",
                            "projects": [],
                        },
                    ),
                    "learnToday": sections.get("learnToday", ""),
                    "preparesFor": sections.get("preparesFor", ""),
                    "connections": sections.get("connections", []),
                    "likelyQuestions": sections.get("likelyQuestions", []),
                    "teachingEmphasis": sections.get("teachingEmphasis", []),
                    "misconceptions": sections.get("misconceptions", []),
                    "oneMinuteRecap": sections.get("oneMinuteRecap", ""),
                }
        modules.append(module)

    add_em_shared_interview_modules(modules, paragraphs)
    return modules


def phase_for(week_key: str, title: str) -> str:
    lower = (week_key + title).lower()
    if week_key == "orientation":
        return "foundation"
    if week_key == "capstone":
        return "capstone"
    if "interview" in lower or week_key in {"W12", "W13", "W14", "W15", "W10", "W11"}:
        if week_key in {"W10", "W11", "W12", "W13", "W14", "W15"}:
            return "interview-prep"
    if "llmops" in lower or week_key.endswith("a"):
        return "ops"
    if week_key in {"W1", "W2", "W3"}:
        return "foundation"
    return "core"


def parse_shared_notes(body: list[str]) -> list[dict]:
    notes = []
    for line in body:
        t = line.strip()
        if not t:
            continue
        if t.startswith("(") and "Shared framing" in t:
            notes.append({"label": "Shared module", "text": t})
        elif t.lower().startswith("track-specific note"):
            notes.append({"label": "Track-specific note", "text": t.split(":", 1)[-1].strip()})
        elif t.lower().startswith("what em learners already know"):
            notes.append({"label": "What EM learners already know", "text": t.split(":", 1)[-1].strip()})
        else:
            notes.append({"label": None, "text": t})
    return notes


def shared_with_for(mod_id: str) -> str:
    mapping = {
        "em-w06a": "swe-w08a",
        "em-w07a": "swe-w09a",
        "em-w10": "swe-w12",
        "em-w11": "swe-w13",
        "em-w12": "swe-w14",
        "em-w13": "swe-w15",
    }
    return mapping.get(mod_id)


def add_em_shared_interview_modules(modules: list[dict], paragraphs: list[str]) -> None:
    """Add EM W10–W13 shared modules from the interview-prep block."""
    text = "\n".join(paragraphs)
    if "EM W10" not in text:
        return
    specs = [
        ("em-w10", "W10", "Agentic Research Systems", "swe-w12", "EM W10 — Agentic Research Systems"),
        ("em-w11", "W11", "Agentic Text-to-SQL", "swe-w13", "EM W11 — Agentic Text-to-SQL"),
        ("em-w12", "W12", "Multi-Agent Coordination & Shared Intelligence", "swe-w14", "EM W12 — Multi-Agent Coordination"),
        ("em-w13", "W13", "Self-Improving Agents", "swe-w15", "EM W13 — Self-Improving Agents"),
    ]
    existing = {m["id"] for m in modules}
    for mod_id, week_key, title, shared, raw_line in specs:
        if mod_id in existing:
            for m in modules:
                if m["id"] == mod_id:
                    m["kind"] = "shared"
                    m["sharedWith"] = shared
            continue
        delta = ""
        for p in paragraphs:
            if p.strip().startswith(f"EM {week_key}"):
                delta = p.strip()
                break
        modules.append(
            {
                "id": mod_id,
                "track": "em",
                "order": len(modules) + 1,
                "code": f"EM Module ({week_key})",
                "title": title,
                "subtitle": "Shared with SWE track",
                "weekKey": week_key,
                "weekLabel": f"Week {week_key[1:]}",
                "weeks": [week_key],
                "phase": "interview-prep",
                "kind": "shared",
                "sharedWith": shared,
                "notes": [
                    {"label": "Viewing as EM", "text": delta},
                    {
                        "label": "What EM learners already know",
                        "text": "Use EM builds from W1–W7 rather than the SWE project list in section 2.",
                    },
                ],
            }
        )


def build_track_groups(track: str, modules: list[dict]) -> list[dict]:
    if track == "swe":
        return [
            {
                "id": "foundation",
                "title": "Foundation & Knowledge",
                "intro": ["Weeks 1–3 establish the agent loop and RAG foundations."],
            },
            {
                "id": "coordination",
                "title": "Coordination & Integration",
                "intro": ["Weeks 4–7 extend into multi-agent systems and retrieval depth."],
            },
            {
                "id": "trust-ops",
                "title": "Trust, Ops & Model Adaptation",
                "intro": ["Weeks 8–9 add observability, safety, fine-tuning, and LLMOps."],
            },
            {
                "id": "capstone-interview",
                "title": "Capstone & Interview Prep",
                "intro": ["Capstone integration and advanced interview deep dives."],
            },
        ]
    if track == "em":
        return [
            {
                "id": "foundation",
                "title": "Agentic fundamentals for engineering leaders",
                "intro": ["Weeks 1–3 cover reasoning agents, RAG, and orchestration."],
            },
            {
                "id": "productization",
                "title": "Operationalizing agentic systems",
                "intro": ["Weeks 4–7 focus on workflows, evaluation, and optimization."],
            },
            {
                "id": "capstone-interview",
                "title": "Capstone & interview deep dives",
                "intro": ["Capstone leadership integration and shared SWE deep dives."],
            },
        ]
    return [
        {
            "id": "foundation",
            "title": "Agentic product foundations",
            "intro": ["Weeks 1–3 introduce agents, RAG, and orchestration from a PM lens."],
        },
        {
            "id": "architecture",
            "title": "Architecture, evaluation & optimization",
            "intro": ["Weeks 4–7 cover product architecture, evaluation, and fine-tuning decisions."],
        },
        {
            "id": "capstone-interview",
            "title": "Capstone & advanced product modules",
            "intro": ["Capstone product integration and advanced PM modules."],
        },
    ]


def build_document(modules: list[dict]) -> dict:
    tracks = []
    for track_id, track_name, short_name, heading in [
        ("swe", "Software Engineers", "SWE", "TRACK 1 — Software Engineers (SWEs)"),
        ("em", "Engineering Managers", "EM", "TRACK 2 — Engineering Managers (EMs)"),
        ("pm", "Product Managers / TPMs", "PM", "TRACK 3 — Product Managers / TPMs (PMs–TPMs)"),
    ]:
        track_modules = [m for m in modules if m["track"] == track_id]
        module_index = {}
        for m in track_modules:
            wk = m["weekKey"]
            if wk:
                module_index[wk] = m["id"]
        tracks.append(
            {
                "id": track_id,
                "name": track_name,
                "shortName": short_name,
                "heading": heading,
                "arc": "",
                "groups": build_track_groups(track_id, track_modules),
                "modules": track_modules,
                "moduleIndex": module_index,
            }
        )

    return {
        "program": "Agentic AI 2.0",
        "documentTitle": "Instructor Context Briefs — Agentic AI 2.0",
        "subtitle": "Cross-module orientation guide for instructors across SWE, EM, and PM/TPM tracks.",
        "version": "2.0",
        "sourceDocument": "Instructor_Context_Briefs_Agentic_AI_2_0_v2.docx",
        "curriculumSheet": {
            "title": "Agentic AI 2.0 Detailed Topic list, Assignment, Pre/Post Class Resources",
            "url": "https://docs.google.com/spreadsheets/",
        },
        "howToUse": [
            "Read the one-minute recap right before class.",
            "Use section 2 to connect prior projects to today's lesson.",
            "Refer to the curriculum sheet for detailed topic coverage.",
        ],
        "conventions": [
            "Python prerequisite is self-paced and not covered in live classes.",
            "Four advanced interview-prep sessions appear at the end of the program.",
            "Pre-class learner emails include setup instructions for each session.",
        ],
        "trackMap": {
            "header": ["Phase", "SWEs (code-first)", "EMs (no-code / n8n)", "PMs–TPMs (no-code / product lens)"],
            "rows": [],
        },
        "trackMapNote": "The three tracks teach the same mental model of agentic AI with track-specific tooling and framing.",
        "sectionTitles": {
            "courseJourney": "1. Course journey",
            "priorKnowledge": "2. What learners already know",
            "learnToday": "3. What learners will learn today",
            "preparesFor": "4. What this prepares learners for",
            "connections": "5. How instructors should connect today's lesson",
            "likelyQuestions": "6. Questions learners are likely to ask",
            "teachingEmphasis": "7. Teaching emphasis",
            "misconceptions": "8. Common misconceptions",
            "oneMinuteRecap": "9. One-minute recap for instructors",
        },
        "tracks": tracks,
    }


def main() -> None:
    doc = Document(str(DOCX_PATH))
    paragraphs = [p.text for p in doc.paragraphs]

    # Track map table
    rows = []
    for row in doc.tables[0].rows:
        rows.append([cell.text.strip() for cell in row.cells])

    modules = find_modules(paragraphs)
    document = build_document(modules)
    document["trackMap"]["rows"] = rows[1:] if len(rows) > 1 else rows

    # Mark known shared modules
    for m in modules:
        sw = shared_with_for(m["id"])
        if sw and m["kind"] != "shared":
            if m["id"] in {"em-w06a", "em-w07a"}:
                m["kind"] = "shared"
                m["sharedWith"] = sw

    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(document, ensure_ascii=False, indent=2), encoding="utf-8")

    index = {
        track["id"]: track["moduleIndex"] for track in document["tracks"]
    }
    OUT_INDEX.write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")

    brief_count = sum(1 for m in modules if m["kind"] == "brief")
    print(f"Wrote {OUT_JSON} ({OUT_JSON.stat().st_size // 1024} KB)")
    print(f"Modules: {len(modules)} (brief={brief_count})")
    print(f"Tracks: {[t['id'] for t in document['tracks']]}")


if __name__ == "__main__":
    main()
